import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.join(os.path.dirname(__file__)))
from app.services.seed_service import STUDENT_SEED_DATA

wb = openpyxl.Workbook()

# Styling definitions
title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
data_font = Font(name="Calibri", size=10)
bold_font = Font(name="Calibri", size=10, bold=True)
center_align = Alignment(horizontal="center", vertical="center")
left_align = Alignment(horizontal="left", vertical="center")

border_thin = Side(border_style="thin", color="CBD5E1")
cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

fill_faculty_title = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Navy Blue
fill_faculty_hdr = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")   # Royal Blue
fill_student_title = PatternFill(start_color="064E3B", end_color="064E3B", fill_type="solid") # Deep Emerald
fill_student_hdr = PatternFill(start_color="059669", end_color="059669", fill_type="solid")   # Emerald
fill_admin_title = PatternFill(start_color="4C1D95", end_color="4C1D95", fill_type="solid")   # Deep Purple
fill_admin_hdr = PatternFill(start_color="6D28D9", end_color="6D28D9", fill_type="solid")     # Purple
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

# -------------------------------------------------------------
# 1. FACULTY SHEET
# -------------------------------------------------------------
ws1 = wb.active
ws1.title = "Faculty Credentials"
ws1.views.sheetView[0].showGridLines = True

ws1.merge_cells("A1:G1")
ws1["A1"] = "NSIT-IFSCS SmartAttend - Faculty Login Credentials"
ws1["A1"].font = title_font
ws1["A1"].fill = fill_faculty_title
ws1["A1"].alignment = center_align
ws1.row_dimensions[1].height = 35

fac_headers = ["Sr No", "Faculty Name", "Employee Code", "Department", "Email (Login ID)", "Password", "Abbrev"]
for col_num, h in enumerate(fac_headers, 1):
    c = ws1.cell(row=2, column=col_num, value=h)
    c.font = header_font
    c.fill = fill_faculty_hdr
    c.alignment = center_align
    c.border = cell_border
ws1.row_dimensions[2].height = 25

faculty_data = [
    (1, "Dr. Akash Thakkar", "FAC001", "Department of Cyber Security", "akash.thakkar@nsit.ac.in", "Faculty@123", "AT"),
    (2, "Dr. Minal Shah", "FAC002", "Department of Computer Science", "minal.shah@nsit.ac.in", "Faculty@123", "MS"),
    (3, "Prof. (Dr.) Sailesh Iyer", "FAC003", "Department of Cyber Security", "sailesh.iyer@nsit.ac.in", "Faculty@123", "SI"),
    (4, "Dr. Nikunj Tahilramani", "FAC004", "Department of Cyber Security", "nikunj.tahilramani@nsit.ac.in", "Faculty@123", "NT"),
    (5, "Dr. Vishali Sharma", "FAC005", "Department of Cyber Security", "vishali.sharma@nsit.ac.in", "Faculty@123", "VS"),
    (6, "Ms. Hepi Suthar", "FAC006", "Department of Cyber Security", "hepi.suthar@nsit.ac.in", "Faculty@123", "HS"),
]

for row_idx, row_data in enumerate(faculty_data, start=3):
    ws1.row_dimensions[row_idx].height = 22
    for col_idx, val in enumerate(row_data, start=1):
        c = ws1.cell(row=row_idx, column=col_idx, value=val)
        c.font = data_font
        c.border = cell_border
        if col_idx in (1, 3, 6, 7):
            c.alignment = center_align
        else:
            c.alignment = left_align
        if row_idx % 2 == 0:
            c.fill = fill_zebra

for col in ws1.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

# -------------------------------------------------------------
# 2. STUDENT SHEET
# -------------------------------------------------------------
ws2 = wb.create_sheet(title="Student Credentials")
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells("A1:F1")
ws2["A1"] = "NSIT-IFSCS SmartAttend - Student Login Credentials (Class Room 109, Sem-III)"
ws2["A1"].font = title_font
ws2["A1"].fill = fill_student_title
ws2["A1"].alignment = center_align
ws2.row_dimensions[1].height = 35

stu_headers = ["Roll No", "Student Name", "Enrollment No (Primary Login)", "Student Email (Alt Login)", "Password", "Class / Branch"]
for col_num, h in enumerate(stu_headers, 1):
    c = ws2.cell(row=2, column=col_num, value=h)
    c.font = header_font
    c.fill = fill_student_hdr
    c.alignment = center_align
    c.border = cell_border
ws2.row_dimensions[2].height = 25

for row_idx, (enroll, roll, name) in enumerate(STUDENT_SEED_DATA, start=3):
    ws2.row_dimensions[row_idx].height = 22
    email = f"student{roll}@nsit.ac.in"
    row_vals = [roll, name, enroll, email, "Student@123", "B.Tech-M.Tech CSE Sem-III"]
    for col_idx, val in enumerate(row_vals, start=1):
        c = ws2.cell(row=row_idx, column=col_idx, value=val)
        c.font = data_font
        c.border = cell_border
        if col_idx in (1, 3, 5):
            c.alignment = center_align
        else:
            c.alignment = left_align
        if row_idx % 2 == 0:
            c.fill = fill_zebra

for col in ws2.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws2.column_dimensions[col_letter].width = max(max_len + 4, 12)

# -------------------------------------------------------------
# 3. ADMIN & PORTAL INFO SHEET
# -------------------------------------------------------------
ws3 = wb.create_sheet(title="Admin & Portals")
ws3.views.sheetView[0].showGridLines = True

ws3.merge_cells("A1:E1")
ws3["A1"] = "NSIT-IFSCS SmartAttend - System Access Portals & Admin Credentials"
ws3["A1"].font = title_font
ws3["A1"].fill = fill_admin_title
ws3["A1"].alignment = center_align
ws3.row_dimensions[1].height = 35

admin_headers = ["Role", "Account Description", "Username / Login ID", "Password", "Live Portal URL"]
for col_num, h in enumerate(admin_headers, 1):
    c = ws3.cell(row=2, column=col_num, value=h)
    c.font = header_font
    c.fill = fill_admin_hdr
    c.alignment = center_align
    c.border = cell_border
ws3.row_dimensions[2].height = 25

admin_data = [
    ("System Admin", "Master Administrator Account", "admin@nsit.ac.in", "Admin@123", "https://nsit-attendance.netlify.app/admin"),
    ("Faculty Member", "Faculty Teaching Portal (6 Faculty)", "Individual Email (e.g. akash.thakkar@nsit.ac.in)", "Faculty@123", "https://nsit-attendance.netlify.app/faculty"),
    ("Student", "Student Attendance Portal (33 Students)", "Enrollment No (e.g. 251943004001)", "Student@123", "https://nsit-attendance.netlify.app/student"),
]

for row_idx, row_vals in enumerate(admin_data, start=3):
    ws3.row_dimensions[row_idx].height = 24
    for col_idx, val in enumerate(row_vals, start=1):
        c = ws3.cell(row=row_idx, column=col_idx, value=val)
        c.font = data_font
        c.border = cell_border
        if col_idx in (1, 4):
            c.alignment = center_align
        else:
            c.alignment = left_align
        if row_idx % 2 == 0:
            c.fill = fill_zebra

for col in ws3.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws3.column_dimensions[col_letter].width = max(max_len + 4, 15)

# Save file to root workspace
out_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "NSIT_SmartAttend_Credentials.xlsx"))
wb.save(out_path)
print(f"Credentials sheet generated successfully at: {out_path}")
